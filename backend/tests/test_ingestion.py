import io
import zipfile
import fitz
import docx
from fastapi.testclient import TestClient
import pytest
from app.main import app
from app.core.database import get_db

client = TestClient(app)


# ==============================================================================
# Helpers to generate synthetic in-memory test documents
# ==============================================================================

def create_synthetic_pdf(text: str) -> bytes:
    """Generate a valid, text-based PDF in-memory using PyMuPDF."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), text, fontsize=11)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_scanned_pdf() -> bytes:
    """Generate a PDF containing an empty page or minimal characters (< 50 chars)."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Scan", fontsize=11)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_synthetic_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    """Generate a valid DOCX file in-memory using python-docx."""
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r_idx, row in enumerate(table_rows):
            for c_idx, cell_text in enumerate(row):
                table.rows[r_idx].cells[c_idx].text = cell_text
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ==============================================================================
# Extraction Tests (PDF & DOCX)
# ==============================================================================

def test_ingest_text_pdf_successful():
    """Verify text-based PDF extracts academic fields, attaches provenance, and leaves missing fields missing."""
    pdf_content = (
        "Sarah Chen\n"
        "University: University of Sao Paulo\n"
        "Degree: Bachelor of Science in Data Science\n"
        "CGPA: 8.8 / 10.0\n"
        "Projects & Experience:\n"
        "- Genomic Transformer: Deep learning model for gene classification\n"
        "https://github.com/sarahchen\n"
    )
    pdf_bytes = create_synthetic_pdf(pdf_content)

    response = client.post(
        "/api/v1/profiles/ingest",
        files={"file": ("resume.pdf", pdf_bytes, "application/pdf")},
        data={"document_type": "resume"},
    )
    assert response.status_code == 200
    data = response.json()

    assert data["status"] in ["success", "partial"]
    assert data["document_type"] == "resume"
    cand = data["candidate_profile"]
    assert cand["name"] == "Sarah Chen"
    assert cand["institution"] == "University of Sao Paulo"
    assert "Data Science" in (cand["field_of_study"] or "")
    assert cand["academic_score"]["score"] == 8.8
    assert cand["academic_score"]["scale"] == 10.0
    assert len(cand["projects_experiences"]) >= 1
    assert "https://github.com/sarahchen" in cand["public_profile_urls"]

    # Verify provenance attached
    assert "name" in data["provenance"]
    assert data["provenance"]["name"]["extracted_value"] == "Sarah Chen"
    assert "academic_score" in data["provenance"]

    # Verify unsupported/unmentioned fields remain missing
    assert "budget_constraint" in data["missing_fields"]
    assert "target_destinations" in data["missing_fields"]
    assert cand["budget_constraint"] is None
    assert cand["target_destinations"] == []


def test_ingest_docx_with_paragraphs_and_tables():
    """Verify DOCX extraction handles paragraphs and tables with provenance."""
    paragraphs = [
        "Elena Rostova",
        "Institution: Technical University of Sofia",
        "Program: Bachelor of Engineering in Robotics & Autonomous Systems",
        "Academic Score: 5.85 / 6.0",
        "Projects & Experience:",
        "- Drone Navigation: Visual-inertial odometry with ROS2",
    ]
    table_data = [
        ["Course", "Grade"],
        ["Autonomous Systems", "6.0 / 6.0"],
    ]
    docx_bytes = create_synthetic_docx(paragraphs, table_data)

    response = client.post(
        "/api/v1/profiles/ingest",
        files={"file": ("transcript.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"document_type": "transcript"},
    )
    assert response.status_code == 200
    data = response.json()

    assert data["document_type"] == "transcript"
    cand = data["candidate_profile"]
    assert cand["name"] == "Elena Rostova"
    assert cand["institution"] == "Technical University of Sofia"
    assert cand["academic_score"]["score"] == 5.85
    assert cand["academic_score"]["scale"] == 6.0
    assert "name" in data["provenance"]
    assert "academic_score" in data["provenance"]


def test_ingest_transcript_without_score_scale():
    """Verify that score without explicit scale does not invent a scale (remains scale=None)."""
    text = (
        "Candidate Name: Marcus Vance\n"
        "College: Manchester College of Technology\n"
        "Degree: B.Tech in Mechanical Engineering\n"
        "CGPA: 7.9\n"
    )
    pdf_bytes = create_synthetic_pdf(text)

    response = client.post(
        "/api/v1/profiles/ingest",
        files={"file": ("transcript.pdf", pdf_bytes, "application/pdf")},
        data={"document_type": "transcript"},
    )
    assert response.status_code == 200
    data = response.json()
    cand = data["candidate_profile"]
    assert cand["name"] == "Marcus Vance"
    assert cand["academic_score"]["score"] == 7.9
    assert cand["academic_score"]["scale"] is None  # Never invented!


def test_ingest_scanned_pdf_returns_ocr_required():
    """Verify scanned or low-text PDF explicitly returns ocr_required status without fabricating profile."""
    scanned_bytes = create_scanned_pdf()

    response = client.post(
        "/api/v1/profiles/ingest",
        files={"file": ("scanned_resume.pdf", scanned_bytes, "application/pdf")},
        data={"document_type": "resume"},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "ocr_required"
    assert data["requires_human_review"] is True
    assert data["candidate_profile"]["name"] is None
    assert len(data["warnings"]) > 0
    assert "OCR" in data["warnings"][0]


def test_ingest_detects_conflicts_without_silently_overwriting():
    """Verify multiple conflicting values are reported in conflicts list."""
    text = (
        "Name: Rahul Sharma\n"
        "Also known as: Rahul Verma\n"
        "University: National Institute of Technology\n"
        "University: Delhi Technological University\n"
        "CGPA: 8.2\n"
        "CGPA: 8.9\n"
    )
    pdf_bytes = create_synthetic_pdf(text)

    response = client.post(
        "/api/v1/profiles/ingest",
        files={"file": ("resume.pdf", pdf_bytes, "application/pdf")},
        data={"document_type": "resume"},
    )
    assert response.status_code == 200
    data = response.json()
    assert len(data["conflicts"]) > 0
    conflict_fields = [c["field_name"] for c in data["conflicts"]]
    assert "academic_score" in conflict_fields or "institution" in conflict_fields


# ==============================================================================
# Security & Upload Boundary Tests
# ==============================================================================

def test_ingest_rejects_unsupported_extension():
    """Verify files without .pdf or .docx are rejected with 400 Bad Request."""
    response = client.post(
        "/api/v1/profiles/ingest",
        files={"file": ("resume.txt", b"Some plain text resume", "text/plain")},
    )
    assert response.status_code == 400
    assert "Unsupported file format" in response.json()["detail"]


def test_ingest_rejects_mismatched_content_signature():
    """Verify file with .pdf extension containing plain text is rejected by magic byte check."""
    fake_pdf = b"This is not a real PDF document!"
    response = client.post(
        "/api/v1/profiles/ingest",
        files={"file": ("fake.pdf", fake_pdf, "application/pdf")},
    )
    assert response.status_code == 400
    assert "does not match detected file content signature" in response.json()["detail"]


def test_ingest_rejects_empty_file():
    """Verify 0-byte file upload is rejected with 400 Bad Request."""
    response = client.post(
        "/api/v1/profiles/ingest",
        files={"file": ("empty.pdf", b"", "application/pdf")},
    )
    assert response.status_code == 400
    assert "empty" in response.json()["detail"]


def test_ingest_rejects_oversized_file():
    """Verify upload exceeding MAX_UPLOAD_SIZE_BYTES is rejected with 413 Payload Too Large."""
    # Create slightly over 10 MiB dummy payload
    oversized = b"%PDF-" + b"0" * (10 * 1024 * 1024 + 1024)
    response = client.post(
        "/api/v1/profiles/ingest",
        files={"file": ("large.pdf", oversized, "application/pdf")},
    )
    assert response.status_code == 413
    assert "maximum allowable size" in response.json()["detail"]


def test_ingest_rejects_malformed_docx():
    """Verify corrupted zip/docx archive returns 400 or 422 error without leaking stack trace."""
    corrupted_docx = b"PK\x03\x04" + b"Corrupted content here"
    response = client.post(
        "/api/v1/profiles/ingest",
        files={"file": ("corrupt.docx", corrupted_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code in [400, 422]
    detail = response.json()["detail"]
    assert "Traceback" not in detail
    assert "C:\\" not in detail


# ==============================================================================
# Confirmation & Persistence Tests
# ==============================================================================

def test_ingestion_does_not_persist_profile_automatically(db_session):
    """Verify that calling /ingest does not persist any record into the database."""
    app.dependency_overrides[get_db] = lambda: db_session
    try:
        pdf_bytes = create_synthetic_pdf("Sarah Chen\nUniversity: MIT\nCGPA: 4.0/4.0")
        response = client.post(
            "/api/v1/profiles/ingest",
            files={"file": ("resume.pdf", pdf_bytes, "application/pdf")},
        )
        assert response.status_code == 200

        # Query database session: student_profiles must remain 0
        from sqlalchemy import select
        from app.models.entities import StudentProfileModel
        profiles = db_session.execute(select(StudentProfileModel)).scalars().all()
        assert len(profiles) == 0
    finally:
        app.dependency_overrides.pop(get_db, None)


def test_confirm_profile_persists_confirmed_data(db_session):
    """Verify that explicitly confirmed profile is persisted via /confirm."""
    app.dependency_overrides[get_db] = lambda: db_session
    try:
        profile_payload = {
            "name": "Sarah Chen",
            "institution": "University of Sao Paulo",
            "degree_program": "B.Sc. Data Science",
            "field_of_study": "AI & Statistics",
            "country": "Brazil",
            "academic_score": {"score": 8.8, "scale": 10.0},
            "target_destinations": ["Germany"],
            "academic_interests": ["Machine Learning"],
            "projects_experiences": [
                {
                    "title": "Genomic Transformer",
                    "summary": "Deep learning model for gene classification",
                }
            ],
            "public_profile_urls": ["https://github.com/sarahchen"],
        }

        response = client.post(
            "/api/v1/profiles/confirm",
            json=profile_payload,
        )
        assert response.status_code == 201
        created = response.json()
        assert "id" in created
        assert created["name"] == "Sarah Chen"
        assert created["country"] == "Brazil"
        assert created["academic_score"]["score"] == 8.8

        # Verify in database
        from app.models.entities import StudentProfileModel
        db_profile = db_session.get(StudentProfileModel, created["id"])
        assert db_profile is not None
        assert db_profile.name == "Sarah Chen"
        assert db_profile.institution == "University of Sao Paulo"
        assert db_profile.academic_score["score"] == 8.8
    finally:
        app.dependency_overrides.pop(get_db, None)


def test_confirm_profile_rejects_invalid_payload():
    """Verify that invalid payload with extra fields is rejected with 422 Unprocessable Entity."""
    response = client.post(
        "/api/v1/profiles/confirm",
        json={"name": "Test", "unauthorized_extra_field": "error"},
    )
    assert response.status_code == 422
